"""
main.py - Excel Copilot backend (built to cope with big workbooks)

Run with:  uvicorn main:app --reload --port 8000

How it scales
  * The uploaded file is saved to disk and only its XML "table of contents" is read at
    upload time, so uploading is instant however big the workbook is.
  * Sheets are read with a fast Rust reader (python-calamine): ~0.5 s for 75,000 rows.
    The first sheet is loaded before the upload call returns, the others on a background thread.
  * Nothing is written into an openpyxl workbook while you work. Every change is recorded in
    `writes` (what should end up in the file) and in `grid` (what to show). That keeps memory low.
  * The grid is paged (default 1,000 rows), duplicates are found on the server, and totals /
    lookups use indexes, so 75,000-row sheets stay responsive.
  * Download is built in the background with progress:
      - small workbooks: the original is opened with openpyxl, so every style is kept;
      - big workbooks:   sheets are streamed out with xlsxwriter (all data and formulas are kept,
                         cell styling is not).
"""
import datetime as dt
import json
import os
import re
import shutil
import tempfile
import threading
import time
import uuid
import zipfile
import xml.etree.ElementTree as ET
from bisect import bisect_left, bisect_right
from collections import defaultdict
from typing import List, Optional, Union

import numpy as np
import openpyxl
import pandas as pd
import xlsxwriter
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from openpyxl.utils import column_index_from_string, get_column_letter
from pydantic import BaseModel
from python_calamine import CalamineWorkbook

from expr_engine import dependency_info, op_expr, referenced_columns, serial
from formula_engine import OPERATIONS as _LEGACY_OPS
from sql_engine import SqlError, Workbench, from_sql_value, sql_value

OPERATIONS = {**_LEGACY_OPS, "EXPR": op_expr}

FRONTEND_VERSION = 17                 # the page files this server expects (bumped whenever the UI changes)
API_VERSION = 6                       # bump when the page needs newer server code (checked by the frontend)
STREAM_CELL_THRESHOLD = 1_000_000     # above this many cells the download is streamed (see module docstring)
PAGE_DEFAULT, PAGE_MAX = 1000, 5000
MAX_SESSIONS = 4
WORKDIR = os.path.join(tempfile.gettempdir(), "excel_copilot")

app = FastAPI(title="Excel Copilot API")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],   # tighten this to your frontend URL when you deploy
    allow_methods=["*"],
    allow_headers=["*"],
)

SESSIONS = {}


# ---------------------------------------------------------------------------
# small helpers
# ---------------------------------------------------------------------------
def _clean(v):
    """Make a cell value JSON-safe."""
    if isinstance(v, (list, dict)):
        return v
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat()
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    if hasattr(v, "item"):          # numpy scalar -> python scalar
        return v.item()
    return v


def _native(v):
    """numpy / pandas scalar -> plain python value."""
    if v is None or (not isinstance(v, (list, dict)) and pd.isna(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if hasattr(v, "item"):
        return v.item()
    return v


def _norm_cell(v):
    """calamine gives "" for empty cells and 5.0 for the number 5 - undo both."""
    if v == "":
        return None
    if v.__class__ is float and v.is_integer() and abs(v) < 1e15:
        return int(v)
    return v


def _display_text(v):
    """How a value reads in a cell / on the clipboard."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat()
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    return str(v)


def _get_session(session_id: str):
    if session_id not in SESSIONS:
        raise HTTPException(404, "Session not found or expired. Please re-upload your file.")
    return SESSIONS[session_id]


def _check_sheet(session, name):
    if name not in session["names"]:
        raise HTTPException(404, "Sheet not found")


# ---------------------------------------------------------------------------
# reading the workbook
# ---------------------------------------------------------------------------
def _read_layout(path):
    """Sheet names (in order), {name: (rows, cols)} and {name: xml part} from the zip's XML - instant for any size."""
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    z = zipfile.ZipFile(path)
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    target = {r.get("Id"): r.get("Target") for r in rels}
    names, dims, parts = [], {}, {}
    for sh in wb.find("m:sheets", ns):
        name = sh.get("name")
        t = target.get(sh.get("{%s}id" % ns["r"]), "")
        part = t.lstrip("/") if t.startswith("/") else "xl/" + t
        names.append(name)
        parts[name] = part
        rows = cols = 0
        try:
            with z.open(part) as f:
                head = f.read(8000).decode("utf8", "ignore")
            m = re.search(r'<dimension ref="([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?"', head)
            if m:
                rows = int(m.group(4) or m.group(2))
                cols = column_index_from_string(m.group(3) or m.group(1))
        except KeyError:
            pass
        dims[name] = (rows, cols)
    return names, dims, parts


def _new_session(path, filename):
    names, dims, parts = _read_layout(path)
    cells = sum(r * c for r, c in dims.values())
    return {
        "path": path, "filename": filename, "names": names, "dims": dims, "parts": parts, "cells": cells,
        "mode": "stream" if cells > STREAM_CELL_THRESHOLD else "full",
        "meta": {n: {"header_row": 1} for n in names},
        "grid": {n: None for n in names},          # raw rows exactly as they will be shown
        "writes": {n: {} for n in names},          # {(row, col): value or "=formula"}   1-based worksheet coordinates
        "wfmt": {n: {} for n in names},            # {(row, col): number format}
        "orig_formulas": {},                       # only kept for small workbooks
        "dfs": {}, "cols": {}, "ver": {n: 0 for n in names},
        "colcache": {}, "dupcache": {}, "dupcolcache": {}, "recipes": [], "errors": {},
        "preloaded": 0, "export": None, "added": [],
        "deleted_rows": {}, "row_map": {},                # sheet -> bookkeeping for rows removed by dedup
        "lock": threading.RLock(), "created": time.time(),
    }


def _orig_to_current(dels, r):
    """A row number in the ORIGINAL file -> its row number now (or None if that row was deleted)."""
    i = bisect_left(dels, r)
    if i < len(dels) and dels[i] == r:
        return None
    return r - i


def _load_orig_formulas(session, name):
    """Read the formulas that were already in the uploaded file. Row numbers are adjusted for any
    rows removed since (by Duplicates), so this always matches the sheet's CURRENT layout."""
    found = {}
    try:
        wb = openpyxl.load_workbook(session["path"], read_only=True, data_only=False)
        for r, row in enumerate(wb[name].iter_rows(values_only=True), start=1):
            for c, v in enumerate(row, start=1):
                if isinstance(v, str) and v.startswith("="):
                    found[(r, c)] = v
        wb.close()
    except Exception:
        pass
    dels = session["deleted_rows"].get(name)
    if dels:
        found = {(nr, c): v for (r, c), v in found.items() if (nr := _orig_to_current(dels, r)) is not None}
    session["orig_formulas"][name] = found


def _ensure_grid(session, name):
    with session["lock"]:
        if session["grid"][name] is not None:
            return
        rows = CalamineWorkbook.from_path(session["path"]).get_sheet_by_name(name).to_python(skip_empty_area=False)
        session["grid"][name] = [[_norm_cell(v) for v in r] for r in rows]
        if session["mode"] == "full":
            _load_orig_formulas(session, name)


def _load_sheet(rows, header_row):
    """DataFrame of the data below the header row + the worksheet column index (1-based) of each column."""
    if header_row - 1 >= len(rows):
        return pd.DataFrame(), []
    width = max((len(r) for r in rows), default=0)
    for r in rows:
        if len(r) < width:
            r.extend([None] * (width - len(r)))
    seen, headers = {}, []
    for i, h in enumerate(rows[header_row - 1]):
        name = f"col_{i}" if h is None else str(h)
        if name in seen:                                   # duplicate header -> name, name_2, name_3 ...
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        headers.append(name)
    data = rows[header_row:]
    end = len(data)
    while end and all(c is None for c in data[end - 1]):
        end -= 1
    df = pd.DataFrame(data[:end], columns=headers) if end else pd.DataFrame(columns=headers)
    keep = [i for i, c in enumerate(df.columns) if not (str(c).startswith("col_") and df.iloc[:, i].isna().all())]
    return df.iloc[:, keep], [i + 1 for i in keep]


def _df(session, name):
    with session["lock"]:
        df = session["dfs"].get(name)
        if df is None:
            _ensure_grid(session, name)
            df, cols = _load_sheet(session["grid"][name], session["meta"][name]["header_row"])
            session["dfs"][name], session["cols"][name] = df, cols
        return df


def _cols(session, name):
    _df(session, name)
    return session["cols"][name]


def _touch(session, name):
    """The sheet's data changed: drop what was derived from it."""
    with session["lock"]:
        session["dfs"].pop(name, None)
        session["ver"][name] += 1
        for k in [k for k in session["colcache"] if k[0] == name]:
            del session["colcache"][k]


def _sheet_ctx(session, name):
    df, cols = _df(session, name), _cols(session, name)
    header_row = session["meta"][name]["header_row"]
    return {
        "sheet_name": name, "header_row": header_row, "last_row": header_row + len(df),
        "columns": {c: cols[i] - 1 for i, c in enumerate(df.columns)},      # name -> worksheet idx0
        "df": df, "ver": session["ver"][name], "_cache": session["colcache"],
    }


class _LazyAll:
    """{sheet: ctx} that only loads a sheet when a formula actually looks at it."""
    def __init__(self, session, current_ctx=None):
        self.s, self.cache = session, {}
        if current_ctx:
            self.cache[current_ctx["sheet_name"]] = current_ctx

    def __contains__(self, name):
        return name in self.s["names"]

    def __getitem__(self, name):
        if name not in self.cache:
            self.cache[name] = _sheet_ctx(self.s, name)
        return self.cache[name]


def _add_sheet(session, name, columns, rows):
    """A brand-new sheet (e.g. saved query results). Lives in memory like any loaded sheet and is written on download."""
    name = re.sub(r"[\[\]:*?/\\]", "_", (name or "Query result").strip())[:31] or "Query result"
    base, k = name, 2
    while name.lower() in (n.lower() for n in session["names"]):
        suffix = f"_{k}"
        name, k = base[:31 - len(suffix)] + suffix, k + 1
    with session["lock"]:
        session["names"].append(name)
        session["grid"][name] = [[str(c) for c in columns]] + [[from_sql_value(v) for v in r] for r in rows]
        session["dims"][name] = (len(rows) + 1, len(columns))
        session["meta"][name] = {"header_row": 1}
        session["writes"][name], session["wfmt"][name], session["ver"][name] = {}, {}, 0
        session["added"].append(name)
    return name


def _cell_formula(session, name, r, c):
    w = session["writes"][name].get((r, c))
    if isinstance(w, str) and w.startswith("="):
        return w
    if (r, c) in session["writes"][name]:              # typed over a formula -> not a formula any more
        return None
    return session["orig_formulas"].get(name, {}).get((r, c))


def _set_cell(session, name, r, c, value, formula=None, fmt=None):
    """r, c: 1-based worksheet coordinates. Updates what is shown AND what will be saved."""
    grid = session["grid"][name]
    width = len(grid[0]) if grid else 0
    while len(grid) < r:
        grid.append([None] * max(width, c))
    row = grid[r - 1]
    if len(row) < c:
        row.extend([None] * (c - len(row)))
    row[c - 1] = value
    session["writes"][name][(r, c)] = formula if formula else value
    if fmt:
        session["wfmt"][name][(r, c)] = fmt


def _delete_rows(session, name, idxs0):
    """Remove data rows (0-based, relative to the sheet's data) for good: from the live grid, from every
    recorded edit/formula, and (at export time) from the original file too. Returns how many rows were removed."""
    with session["lock"]:
        header = session["meta"][name]["header_row"]
        grid = session["grid"][name]
        wrows = sorted({header + 1 + i for i in idxs0 if 0 <= header + i < len(grid)})
        if not wrows:
            return 0
        wrowset = set(wrows)
        rm = session["row_map"].setdefault(name, list(range(1, len(grid) + 1)))
        orig_removed = [rm[wr - 1] for wr in wrows if wr - 1 < len(rm)]
        for wr in reversed(wrows):
            del grid[wr - 1]
            if wr - 1 < len(rm):
                del rm[wr - 1]
        dels = session["deleted_rows"].setdefault(name, [])
        dels.extend(orig_removed)
        dels.sort()

        def shift(r):
            return r - bisect_right(wrows, r)

        for store_name in ("writes", "wfmt"):
            store = session[store_name][name]
            new = {(shift(r), c): v for (r, c), v in store.items() if r not in wrowset}
            store.clear()
            store.update(new)
        of = session["orig_formulas"].get(name)
        if of:
            session["orig_formulas"][name] = {(shift(r), c): v for (r, c), v in of.items() if r not in wrowset}
        rows_dim, cols_dim = session["dims"].get(name, (0, 0))
        session["dims"][name] = (max(0, rows_dim - len(wrows)), cols_dim)
        session["dupcache"] = {k: v for k, v in session["dupcache"].items() if k[0] != name}
        session["dupcolcache"] = {k: v for k, v in session["dupcolcache"].items() if k[0] != name}
        _touch(session, name)
        return len(wrows)


def _reapply_recipe(session, rec):
    """Re-run a saved Formula-Builder step over every row of its sheet: needed after rows move (Duplicates),
    since a row-local formula's own row references would otherwise still point at its old position."""
    try:
        _, ctx, formula_fn, values = _build(rec["body"])
    except Exception:
        return
    name, out_col, header_row = rec["sheet"], rec["out_col"], ctx["header_row"]
    for i, v in enumerate(_native(v) for v in values):
        r = header_row + 1 + i
        _set_cell(session, name, r, out_col, v, formula_fn(r), "yyyy-mm-dd" if isinstance(v, (dt.date, dt.datetime)) else None)
    _touch(session, name)


# ---------------------------------------------------------------------------
# duplicates (server side, so it works on every row, not just the visible page)
# ---------------------------------------------------------------------------
def _dup_columns(session, name, cols, norm):
    """Each ticked column, as normalised comparison text (cached - several screens share this)."""
    ver = session["ver"][name]
    key = (name, ver, tuple(cols), bool(norm))
    cache = session["dupcolcache"]
    if key in cache:
        return cache[key]
    df = _df(session, name)
    out = []
    for c in cols:
        col = []
        for v in df.iloc[:, c].tolist():
            t = "" if v is None or v != v else _display_text(v)
            col.append(" ".join(t.split()).lower() if norm else t)
        out.append(col)
    if len(cache) > 8:
        cache.clear()
    cache[key] = out
    return out


def _dup_cols(session, name, spec):
    cols = [c for c in spec.get("cols", []) if 0 <= c < _df(session, name).shape[1]]
    if not cols:
        raise HTTPException(400, "Tick at least one column to check for duplicates.")
    return cols


def _dup_info(session, name, spec):
    cols = _dup_cols(session, name, spec)
    norm = spec.get("norm", True)
    skip_empty = spec.get("skip_empty", True)
    ver = session["ver"][name]
    key = (name, ver, tuple(cols), bool(norm), bool(skip_empty))
    cache = session["dupcache"]
    if key in cache:
        return cache[key]
    columns = _dup_columns(session, name, cols, norm)
    n = len(columns[0]) if columns else 0
    groups = defaultdict(list)
    for i, k in enumerate(zip(*columns)):
        if skip_empty and all(p == "" for p in k):
            continue
        groups[k].append(i)
    gid = np.full(n, -1, dtype=np.int32)
    size = np.zeros(n, dtype=np.int32)
    group_keys, g = [], 0
    for k, lst in groups.items():
        if len(lst) > 1:
            gid[lst], size[lst] = g, len(lst)
            group_keys.append(k)
            g += 1
    info = {"gid": gid, "size": size, "groups": g, "rows": int((gid >= 0).sum()), "checked": n,
            "dup_ids": np.nonzero(gid >= 0)[0], "group_keys": group_keys, "cols": cols}
    if len(cache) > 6:
        cache.clear()
    cache[key] = info
    return info


def _dup_delete_idx(dup, keep):
    """0-based row indices to remove so every duplicate group is left with exactly one row."""
    gid = dup["gid"]
    keepers = {}
    order = range(len(gid)) if keep != "last" else range(len(gid) - 1, -1, -1)
    for i in order:
        g = int(gid[i])
        if g < 0:
            continue
        keepers.setdefault(g, i)
    keep_set = set(keepers.values())
    return [int(i) for i in dup["dup_ids"].tolist() if i not in keep_set]


# ---------------------------------------------------------------------------
# paging
# ---------------------------------------------------------------------------
def _rows_payload(session, name, idx):
    """Full rows (every column) for an explicit list of 0-based data-row indices - used for paging and
    for the Duplicates screen's "find a specific value" and group-preview lookups."""
    df, cols = _df(session, name), _cols(session, name)
    header_row = session["meta"][name]["header_row"]
    part = df.iloc[idx]
    return {
        "columns": list(df.columns),
        "col_letters": [get_column_letter(c) for c in cols],
        "rows": [[_clean(v) for v in row] for row in part.itertuples(index=False)],
        "formulas": [[_cell_formula(session, name, header_row + 1 + i, wc) for wc in cols] for i in idx],
        "ids": list(idx),
        "header_row": header_row,
    }



# ---------------------------------------------------------------------------
# column filter / sort (a "view" of the sheet: which rows, in what order - the data itself never changes)
# ---------------------------------------------------------------------------
_TEXT_OPS = {"contains", "notcontains", "starts", "ends", "equals", "empty", "notempty"}
_NUM_OPS = {"=", "!=", ">", ">=", "<", "<=", "empty", "notempty"}
_DATE_OPS = {"on", "before", "after", "empty", "notempty"}


def _col_kind(s):
    """'number' | 'date' | 'text' - judged from the first values of the column."""
    vals = s.head(2000).dropna()
    if not len(vals):
        return "text"
    vals = [v for v in vals.iloc[:500].tolist() if v != ""]
    if not vals:
        return "text"
    if all(isinstance(v, (int, float, np.integer, np.floating)) and not isinstance(v, (bool, np.bool_)) for v in vals):
        return "number"
    if all(isinstance(v, (dt.date, np.datetime64)) for v in vals):
        return "date"
    return "text"


def _filter_mask(s, kind, op, value):
    text = s.map(lambda v: _display_text(_native(v)))
    if op == "empty":
        return (text.str.strip() == "").to_numpy()
    if op == "notempty":
        return (text.str.strip() != "").to_numpy()
    value = "" if value is None else str(value)
    if kind == "number":
        if op not in _NUM_OPS:
            raise HTTPException(400, f"“{op}” can't be used on a number column.")
        try:
            x = float(value.replace(",", "").strip())
        except ValueError:
            raise HTTPException(400, f"“{value}” isn't a number.")
        n = pd.to_numeric(s, errors="coerce")
        return {"=": n == x, "!=": n != x, ">": n > x, ">=": n >= x, "<": n < x, "<=": n <= x}[op].fillna(False).to_numpy()
    if kind == "date":
        if op not in _DATE_OPS:
            raise HTTPException(400, f"“{op}” can't be used on a date column.")
        d = pd.to_datetime(value, errors="coerce")
        if pd.isna(d):
            raise HTTPException(400, f"“{value}” isn't a date (try 2024-03-31).")
        col = pd.to_datetime(s, errors="coerce").dt.normalize()
        d = d.normalize()
        return {"on": col == d, "before": col < d, "after": col > d}[op].fillna(False).to_numpy()
    if op not in _TEXT_OPS:
        raise HTTPException(400, f"“{op}” can't be used on a text column.")
    lo, v = text.str.lower(), value.lower()
    if op == "contains":
        return lo.str.contains(v, regex=False).to_numpy()
    if op == "notcontains":
        return (~lo.str.contains(v, regex=False)).to_numpy()
    if op == "starts":
        return lo.str.startswith(v).to_numpy()
    if op == "ends":
        return lo.str.endswith(v).to_numpy()
    return (lo == v).to_numpy()


def _sort_order(s, kind, ids, desc):
    """`ids` (row positions) put in order of the column; empty cells always go last."""
    part = s.iloc[ids]
    text = part.map(lambda v: _display_text(_native(v)))
    blank = (text.str.strip() == "").to_numpy()
    if kind == "number":
        key = pd.to_numeric(part, errors="coerce").to_numpy(dtype=float)
    elif kind == "date":
        key = pd.to_datetime(part, errors="coerce").to_numpy()
    else:
        key = text.str.lower().to_numpy()
    bad = blank | (pd.isna(key) if kind != "text" else False)
    good, rest = ids[~bad], ids[bad]
    gk = key[~bad]
    order = np.argsort(gk, kind="stable")
    if desc:
        order = order[::-1]
    return np.concatenate([good[order], rest])


def _view_ids(session, name, df, view):
    """Row positions that pass every filter, in the requested order (None = the whole sheet as it is)."""
    if not view:
        return None
    filters = [f for f in (view.get("filters") or []) if f]
    sort = view.get("sort")
    if not filters and not sort:
        return None
    key = (name, session["ver"][name], json.dumps([filters, sort], sort_keys=True, default=str))
    cache = session.setdefault("viewcache", {})
    if key in cache:
        return cache[key]
    ncol = df.shape[1]
    mask = np.ones(len(df), dtype=bool)
    for f in filters:
        c = int(f.get("col", -1))
        if not 0 <= c < ncol:
            raise HTTPException(400, "That column no longer exists — clear the filter and try again.")
        s = df.iloc[:, c]
        mask &= _filter_mask(s, _col_kind(s), f.get("op"), f.get("value"))
    ids = np.nonzero(mask)[0]
    if sort:
        c = int(sort.get("col", -1))
        if not 0 <= c < ncol:
            raise HTTPException(400, "That column no longer exists — clear the sort and try again.")
        s = df.iloc[:, c]
        ids = _sort_order(s, _col_kind(s), ids, sort.get("dir") == "desc")
    if len(cache) > 6:
        cache.clear()
    cache[key] = ids
    return ids


def _grid_payload(session, name, offset=0, limit=PAGE_DEFAULT, dup_spec=None, view=None):
    df = _df(session, name)
    total = len(df)
    dup, ids = None, None
    if dup_spec:
        dup = _dup_info(session, name, dup_spec)
        if dup_spec.get("only"):
            ids = dup["dup_ids"]
    vids = _view_ids(session, name, df, view)
    if vids is not None:
        ids = vids if ids is None else vids[np.isin(vids, ids)]      # keeps the sort order
    matched = total if ids is None else len(ids)
    limit = max(1, min(int(limit), PAGE_MAX))
    offset = max(0, int(offset))
    if ids is None:
        lo = min(offset, total)
        idx = list(range(lo, min(lo + limit, total)))
    else:
        idx = [int(i) for i in ids[offset:offset + limit]]
    payload = _rows_payload(session, name, idx)
    payload.update({
        "total_rows": total, "matched_rows": matched, "offset": offset, "limit": limit,
        "col_kinds": [_col_kind(df.iloc[:, c]) for c in range(df.shape[1])],
        "dup": [([int(dup["gid"][i]), int(dup["size"][i])] if dup["gid"][i] >= 0 else None) for i in idx] if dup else None,
        "dup_summary": {"groups": dup["groups"], "rows": dup["rows"], "checked": dup["checked"]} if dup else None,
    })
    return payload


def _sheet_list(session):
    out = []
    for n in session["names"]:
        df = session["dfs"].get(n)
        header = session["meta"][n]["header_row"]
        out.append({
            "name": n, "header_row": header, "loaded": df is not None,
            "columns": list(df.columns) if df is not None else [],
            "row_count": len(df) if df is not None else max(0, session["dims"][n][0] - header),
        })
    return out


def _preload(session):
    for n in session["names"]:
        try:
            _df(session, n)
        except Exception as e:                       # a broken sheet must not stop the others
            session["errors"][n] = str(e)
        session["preloaded"] += 1


def _forget(session):
    for p in (session.get("path"), (session.get("export") or {}).get("path")):
        try:
            if p and os.path.exists(p):
                os.remove(p)
        except OSError:
            pass


# ---------------------------------------------------------------------------
# upload
# ---------------------------------------------------------------------------
def _register_upload(path, filename):
    """Turn a saved workbook (already inside WORKDIR) into a session and return the upload response."""
    try:
        session = _new_session(path, filename)
        _df(session, session["names"][0])                # the sheet you land on is ready straight away
    except Exception as e:
        os.remove(path)
        raise HTTPException(400, f"Couldn't read this workbook: {e}")
    session_id = str(uuid.uuid4())
    while len(SESSIONS) >= MAX_SESSIONS:                 # forget the oldest session (and its temp files)
        oldest = min(SESSIONS, key=lambda k: SESSIONS[k]["created"])
        _forget(SESSIONS.pop(oldest))
    SESSIONS[session_id] = session
    session["preloaded"] = 1
    threading.Thread(target=_preload, args=(session,), daemon=True).start()
    return {"session_id": session_id, "filename": filename, "mode": session["mode"],
            "cells": session["cells"], "sheets": _sheet_list(session)}


@app.post("/api/upload")
def upload(file: UploadFile = File(...)):
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Please upload an .xlsx file (older .xls files are not supported).")
    os.makedirs(WORKDIR, exist_ok=True)
    path = os.path.join(WORKDIR, f"{uuid.uuid4()}.xlsx")
    with open(path, "wb") as out:
        shutil.copyfileobj(file.file, out, 1 << 20)
    return _register_upload(path, file.filename)


@app.get("/api/sheets/{session_id}")
def list_sheets(session_id: str):
    session = _get_session(session_id)
    total = len(session["names"])
    loaded = sum(1 for n in session["names"] if n in session["dfs"])
    return {"sheets": _sheet_list(session), "loaded": loaded, "total": total, "ready": loaded + len(session["errors"]) >= total,
            "errors": session["errors"], "mode": session["mode"]}


# ---------------------------------------------------------------------------
# header row / paging
# ---------------------------------------------------------------------------
class HeaderRowUpdate(BaseModel):
    session_id: str
    sheet_name: str
    header_row: int


@app.post("/api/sheet/header-row")
def set_header_row(body: HeaderRowUpdate):
    session = _get_session(body.session_id)
    _check_sheet(session, body.sheet_name)
    with session["lock"]:
        session["meta"][body.sheet_name]["header_row"] = max(1, body.header_row)
        session["recipes"] = [r for r in session["recipes"] if r["sheet"] != body.sheet_name]
        _touch(session, body.sheet_name)
        df = _df(session, body.sheet_name)
    return {"columns": list(df.columns), "row_count": len(df)}


class PageRequest(BaseModel):
    session_id: str
    sheet_name: str
    offset: int = 0
    limit: int = PAGE_DEFAULT
    dup: Optional[dict] = None            # {"cols": [idx...], "norm": bool, "skip_empty": bool, "only": bool}
    view: Optional[dict] = None           # {"filters": [{"col", "op", "value"}], "sort": {"col", "dir": "asc"|"desc"}}


@app.post("/api/sheet/page")
def get_page(body: PageRequest):
    session = _get_session(body.session_id)
    _check_sheet(session, body.sheet_name)
    return _grid_payload(session, body.sheet_name, body.offset, body.limit, body.dup, body.view)


@app.get("/api/sheet/{session_id}/{sheet_name}")
def get_sheet(session_id: str, sheet_name: str, offset: int = 0, limit: int = PAGE_DEFAULT):
    session = _get_session(session_id)
    _check_sheet(session, sheet_name)
    return _grid_payload(session, sheet_name, offset, limit)


class RangeRequest(BaseModel):
    session_id: str
    sheet_name: str
    r0: int = 0
    r1: int = -1                # inclusive, -1 = last row
    c0: int = 0
    c1: int = -1
    headers: bool = False


def _tsv(s):
    return '"' + s.replace('"', '""') + '"' if re.search(r'[\t\n\r"]', s) else s


@app.post("/api/sheet/range-text")
def range_text(body: RangeRequest):
    """Tab-separated text for a block of cells - used to copy whole columns / sheets (more than one page)."""
    session = _get_session(body.session_id)
    _check_sheet(session, body.sheet_name)
    df = _df(session, body.sheet_name)
    r1 = len(df) - 1 if body.r1 < 0 else min(body.r1, len(df) - 1)
    c1 = df.shape[1] - 1 if body.c1 < 0 else min(body.c1, df.shape[1] - 1)
    if (r1 - body.r0 + 1) * (c1 - body.c0 + 1) > 3_000_000:
        raise HTTPException(400, "That is too many cells to copy at once (limit 3 million).")
    lines = []
    if body.headers:
        lines.append("\t".join(_tsv(str(c)) for c in df.columns[body.c0:c1 + 1]))
    part = df.iloc[body.r0:r1 + 1, body.c0:c1 + 1]
    for row in part.itertuples(index=False):
        lines.append("\t".join(_tsv(_display_text(_native(v))) for v in row))
    return {"text": "\n".join(lines), "cells": max(0, (r1 - body.r0 + 1)) * (c1 - body.c0 + 1)}


# ---------------------------------------------------------------------------
# Duplicates - find, browse, remove (server side, so it works on every row of a 75,000-row sheet)
# ---------------------------------------------------------------------------
class DupSpec(BaseModel):
    session_id: str
    sheet_name: str
    cols: List[int]                # which columns make up the "key" that defines a duplicate
    norm: bool = True              # ignore upper/lower case and extra spaces
    skip_empty: bool = True        # rows where every key column is empty don't count as duplicates of each other


@app.post("/api/dup/summary")
def dup_summary(body: DupSpec):
    """The overview: total records, how many are duplicates, how many would be left."""
    session = _get_session(body.session_id)
    _check_sheet(session, body.sheet_name)
    with session["lock"]:
        dup = _dup_info(session, body.sheet_name, body.dict())
        total = len(_df(session, body.sheet_name))
        extra = dup["rows"] - dup["groups"]          # rows that would be removed, keeping one per group
        return {
            "total_rows": total, "checked": dup["checked"], "key_columns": len(dup["cols"]),
            "groups": dup["groups"], "duplicate_rows": dup["rows"], "extra_rows": extra,
            "unique_rows": total - extra,
        }


class DupGroupsRequest(DupSpec):
    offset: int = 0
    limit: int = 50
    search: str = ""               # only groups whose key contains this text


@app.post("/api/dup/groups")
def dup_groups(body: DupGroupsRequest):
    """A paged, searchable list of the duplicate groups (biggest first), each with its key values."""
    session = _get_session(body.session_id)
    _check_sheet(session, body.sheet_name)
    with session["lock"]:
        dup = _dup_info(session, body.sheet_name, body.dict())
        df = _df(session, body.sheet_name)
        cols = dup["cols"]
        col_names = [str(df.columns[c]) for c in cols]
        first_of = {}
        for i in dup["dup_ids"].tolist():
            first_of.setdefault(int(dup["gid"][i]), int(i))
        q = body.search.strip().lower()
        items = []
        for g in range(dup["groups"]):
            i = first_of[g]
            key_vals = [_display_text(_native(v)) for v in df.iloc[i, cols].tolist()]
            if q and not any(q in kv.lower() for kv in key_vals):
                continue
            items.append({"group": g, "size": int(dup["size"][i]), "key": key_vals, "sample_row": i})
        items.sort(key=lambda x: (-x["size"], x["group"]))
        total = len(items)
        limit = max(1, min(body.limit, 500))
        page = items[body.offset:body.offset + limit]
        return {"key_columns": col_names, "total_groups": total, "groups_with_duplicates": dup["groups"], "items": page}


class DupLookupRequest(DupSpec):
    values: List[str]              # one value per ticked column, in the same order as `cols`


@app.post("/api/dup/lookup")
def dup_lookup(body: DupLookupRequest):
    """Every row whose key exactly matches the given value(s) - "does THIS employee id have duplicates?"."""
    session = _get_session(body.session_id)
    _check_sheet(session, body.sheet_name)
    with session["lock"]:
        cols = _dup_cols(session, body.sheet_name, body.dict())
        if len(body.values) != len(cols):
            raise HTTPException(400, "Give a value for every key column.")
        norm_vals = []
        for v in body.values:
            t = " ".join(str(v).split())
            norm_vals.append(t.lower() if body.norm else t)
        if not any(norm_vals):
            raise HTTPException(400, "Type a value to search for.")
        columns = _dup_columns(session, body.sheet_name, cols, body.norm)
        n = len(columns[0]) if columns else 0
        idx = [i for i in range(n) if all(columns[k][i] == norm_vals[k] for k in range(len(cols)))]
        payload = _rows_payload(session, body.sheet_name, idx)
        dup = _dup_info(session, body.sheet_name, body.dict())
        payload["matches"] = len(idx)
        payload["is_duplicate"] = bool(idx) and int(dup["gid"][idx[0]]) >= 0
        return payload


class DupApplyRequest(DupSpec):
    keep: str = "first"            # first | last - which occurrence survives in each group
    target: str = "existing"       # existing (delete from this sheet) | new (write a new sheet)
    new_sheet_name: str = ""
    content: str = "unique"        # new-sheet content: unique | duplicates | groups


@app.post("/api/dup/apply")
def dup_apply(body: DupApplyRequest):
    session = _get_session(body.session_id)
    _check_sheet(session, body.sheet_name)
    with session["lock"]:
        dup = _dup_info(session, body.sheet_name, body.dict())
        delete_idx = _dup_delete_idx(dup, body.keep)
        if body.target == "existing":
            if not delete_idx:
                return {"mode": "existing", "removed": 0, "total_rows": len(_df(session, body.sheet_name)), "sheet": body.sheet_name}
            name = body.sheet_name
            removed = _delete_rows(session, name, delete_idx)
            for rec in [r for r in session["recipes"] if r["sheet"] == name]:
                _reapply_recipe(session, rec)
            df = _df(session, name)
            return {"mode": "existing", "removed": removed, "total_rows": len(df), "sheet": name,
                     "sheets": _sheet_list(session)}
        df = _df(session, body.sheet_name)
        if body.content == "duplicates":
            idx = delete_idx
        elif body.content == "groups":
            idx = [int(i) for i in dup["dup_ids"].tolist()]
        else:
            drop = set(delete_idx)
            idx = [i for i in range(len(df)) if i not in drop]
        if len(idx) * max(1, df.shape[1]) > 3_000_000:
            raise HTTPException(400, "That result is too big for a new sheet (limit 3 million cells).")
        rows = [[_native(v) for v in r] for r in df.iloc[idx].itertuples(index=False)]
        default_name = f"{body.sheet_name} - " + {"unique": "unique", "duplicates": "duplicates", "groups": "duplicate groups"}[body.content]
        new_name = _add_sheet(session, body.new_sheet_name or default_name, list(df.columns), rows)
        _df(session, new_name)
        return {"mode": "new", "sheet": new_name, "rows": len(rows), "sheets": _sheet_list(session)}


class DupExportRequest(DupSpec):
    keep: str = "first"
    which: str = "unique"          # unique | duplicates | groups
    headers: bool = True


@app.post("/api/dup/export-text")
def dup_export_text(body: DupExportRequest):
    """Tab-separated text of the chosen row set, ready to paste straight into Excel."""
    session = _get_session(body.session_id)
    _check_sheet(session, body.sheet_name)
    with session["lock"]:
        dup = _dup_info(session, body.sheet_name, body.dict())
        delete_idx = _dup_delete_idx(dup, body.keep)
        df = _df(session, body.sheet_name)
        if body.which == "duplicates":
            idx = delete_idx
        elif body.which == "groups":
            idx = [int(i) for i in dup["dup_ids"].tolist()]
        else:
            drop = set(delete_idx)
            idx = [i for i in range(len(df)) if i not in drop]
        if len(idx) * max(1, df.shape[1]) > 3_000_000:
            raise HTTPException(400, "That is too many cells to copy at once (limit 3 million).")
        lines = []
        if body.headers:
            lines.append("\t".join(_tsv(str(c)) for c in df.columns))
        for row in df.iloc[idx].itertuples(index=False):
            lines.append("\t".join(_tsv(_display_text(_native(v))) for v in row))
        return {"text": "\n".join(lines), "rows": len(idx)}


class DeleteRowsRequest(BaseModel):
    session_id: str
    sheet_name: str
    rows: List[int]         # 0-based data-row indices, e.g. the "ids" a page or a dedup lookup handed back


@app.post("/api/sheet/delete-rows")
def sheet_delete_rows(body: DeleteRowsRequest):
    """Delete specific rows, picked by hand - e.g. from the Duplicates screen, where two records can look
    identical in the ticked key columns but differ everywhere else, so only a person can say which to keep."""
    session = _get_session(body.session_id)
    _check_sheet(session, body.sheet_name)
    with session["lock"]:
        name = body.sheet_name
        idxs = sorted({int(i) for i in body.rows if i >= 0})
        if not idxs:
            raise HTTPException(400, "Pick at least one row to delete.")
        removed = _delete_rows(session, name, idxs)
        for rec in [r for r in session["recipes"] if r["sheet"] == name]:
            _reapply_recipe(session, rec)
        df = _df(session, name)
        return {"removed": removed, "total_rows": len(df), "sheet": name, "sheets": _sheet_list(session)}


# ---------------------------------------------------------------------------
# operations (formulas built with the decision-tree builder)
# ---------------------------------------------------------------------------
class ApplyOperation(BaseModel):
    session_id: str
    sheet_name: str
    operation: str            # EXPR (decision-tree builder) | legacy: CONCAT, LOOKUP, ...
    params: dict
    output_column: str = ""
    target_sheet: Optional[str] = None   # legacy LOOKUP only
    label: Optional[str] = None


def _build(body: ApplyOperation, rows=None, preview=False):
    """Compile an operation -> (session, ctx, formula_fn, values). Raises HTTP 400
    for anything the user can fix (missing column, empty box, ...).
    rows: only evaluate these data rows;  preview: only the first few rows."""
    session = _get_session(body.session_id)
    _check_sheet(session, body.sheet_name)
    if body.operation not in OPERATIONS:
        raise HTTPException(400, f"Unknown operation {body.operation}")
    ctx = _sheet_ctx(session, body.sheet_name)
    ctx["_all"] = _LazyAll(session, ctx)
    target_ctx = None
    if body.target_sheet:
        _check_sheet(session, body.target_sheet)
        target_ctx = ctx["_all"][body.target_sheet]
    if preview:
        rows = list(range(min(len(ctx["df"]), 20)))
    try:
        if body.operation == "EXPR":
            formula_fn, values = op_expr(body.params, ctx, target_ctx, rows=rows)
        else:
            formula_fn, values = OPERATIONS[body.operation](body.params, ctx, target_ctx)
            if rows is not None:
                values = [values[i] for i in rows if i < len(values)]
        if len(ctx["df"]):
            formula_fn(ctx["header_row"] + 1)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, str(e) or e.__class__.__name__)
    return session, ctx, formula_fn, values


@app.post("/api/operations/preview")
def preview_operation(body: ApplyOperation):
    """Same as apply, but writes nothing: returns a sample formula and the first values."""
    _, ctx, formula_fn, values = _build(body, preview=True)
    header_row = ctx["header_row"]
    values = list(values)
    return {
        "formula": formula_fn(header_row + 1) if values else None,
        "rows": [{"row": header_row + 1 + i, "value": _clean(v)} for i, v in enumerate(values[:8])],
    }


@app.post("/api/operations/apply")
def apply_operation(body: ApplyOperation):
    session = _get_session(body.session_id)
    with session["lock"]:
        session, ctx, formula_fn, values = _build(body)
        if not body.output_column.strip():
            raise HTTPException(400, "Choose where to put the result")
        name = body.sheet_name
        header_row = ctx["header_row"]
        df, cols = ctx["df"], _cols(session, name)
        grid = session["grid"][name]

        # Destination: an existing column (overwrite) or a brand-new one at the end
        if body.output_column in df.columns:
            out_col = cols[list(df.columns).index(body.output_column)]
        else:
            out_col = max([len(grid[0]) if grid else 0] + cols) + 1

        # Writing a formula into a column it also reads from would be circular
        # (e.g. TRIM column B into column B), so in that case store the values.
        referenced = set()
        if body.operation == "EXPR":
            referenced |= referenced_columns(body.params.get("expr"), name)
        for k in ("columns", "column", "lookup_column", "condition_column", "criteria_column", "value_column"):
            v = body.params.get(k)
            referenced.update(v if isinstance(v, list) else [v])
        static = body.output_column in referenced

        vals = [_native(v) for v in list(values)[:len(df)]]
        _set_cell(session, name, header_row, out_col, body.output_column)
        for i, v in enumerate(vals):
            r = header_row + 1 + i
            is_date = isinstance(v, (dt.date, dt.datetime))
            _set_cell(session, name, r, out_col, v, None if static else formula_fn(r), "yyyy-mm-dd" if is_date else None)
        sample = None if static or not vals else formula_fn(header_row + 1)
        _touch(session, name)

        # keep the "recipe" so later cell edits can refresh this column's values
        session["recipes"] = [r for r in session["recipes"] if (r["sheet"], r["out_col"]) != (name, out_col)]
        if not static:
            info = dependency_info(body.params.get("expr"), name) if body.operation == "EXPR" else {"global": True, "sheets": {name}}
            session["recipes"].append({"body": body, "sheet": name, "out_col": out_col, "info": info})

        df = _df(session, name)
        return {"columns": list(df.columns), "total_rows": len(df), "sample_formula": sample,
                "output_column": body.output_column, "static": static}


# ---------------------------------------------------------------------------
# editing cells
# ---------------------------------------------------------------------------
class CellEdit(BaseModel):
    row: int                                            # 0-based data row (absolute, not page-relative)
    col: int                                            # 0-based column in the grid
    value: Optional[Union[str, int, float, bool]] = None
    formula: Optional[str] = None                       # only used to restore a formula (undo)


class EditBatch(BaseModel):
    session_id: str
    sheet_name: str
    edits: List[CellEdit]


def _parse_input(v):
    """What the user typed / pasted -> a value to store (numbers, TRUE/FALSE, ISO dates, else text)."""
    if v is None or isinstance(v, (bool, int, float)):
        return v
    s = str(v)
    t = s.strip()
    if t == "":
        return None
    if s.startswith("="):
        raise ValueError("Type values in the grid. To create formulas use the Formula Builder on the left.")
    if re.fullmatch(r"-?(0|[1-9]\d*)(\.\d+)?", t):
        return float(t) if "." in t else int(t)
    if t.upper() in ("TRUE", "FALSE"):
        return t.upper() == "TRUE"
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", t)
    if m:
        try:
            return dt.date(int(m[1]), int(m[2]), int(m[3]))
        except ValueError:
            pass
    return s


def _recompute(session, edited, log, log_sheet):
    """Refresh the formula columns we wrote after cells changed, like Excel would.
    Row-local formulas are recalculated for the edited rows only; ones that look at other rows
    (lookups, totals) for the whole column."""
    for rec in list(session["recipes"]):
        info, sheet = rec["info"], rec["sheet"]
        if not (info["sheets"] & set(edited)):
            continue
        rows = None
        if not info["global"] and sheet in edited:
            n_rows = len(_df(session, sheet))
            rows = sorted(i for i in edited[sheet] if i < n_rows)
            if not rows:
                continue
        try:
            _, ctx, _, values = _build(rec["body"], rows=rows)
        except Exception:
            continue                                     # e.g. the header row changed since
        writes, grid, out_col, header = session["writes"][sheet], session["grid"][sheet], rec["out_col"], ctx["header_row"]
        changed = False
        for i, v in zip(range(len(values)) if rows is None else rows, values):
            r = header + 1 + i
            cur = writes.get((r, out_col))
            if isinstance(cur, str) and cur.startswith("="):          # skip cells the user typed over
                v = _native(v)
                if r <= len(grid) and out_col <= len(grid[r - 1]) and grid[r - 1][out_col - 1] != v:
                    grid[r - 1][out_col - 1] = v
                    changed = True
                    if sheet == log_sheet:
                        log.append((i, out_col))
        if changed:
            _touch(session, sheet)


@app.post("/api/sheet/edit")
def edit_cells(body: EditBatch):
    session = _get_session(body.session_id)
    _check_sheet(session, body.sheet_name)
    name = body.sheet_name
    try:
        parsed = [(e, None if e.formula else _parse_input(e.value)) for e in body.edits]
    except ValueError as ex:
        raise HTTPException(400, str(ex))
    with session["lock"]:
        header_row = session["meta"][name]["header_row"]
        df, cols = _df(session, name), _cols(session, name)
        n_before, names_before = len(df), list(df.columns)
        log, edited = [], set()
        for e, value in parsed:
            if e.row < 0 or e.col < 0 or e.col >= len(cols) or e.row > n_before + 2000:
                continue                                 # outside the sheet: ignored (clipped)
            r, c = header_row + 1 + e.row, cols[e.col]
            if e.formula:
                shown = _native(_parse_input(e.value)) if e.value is not None else None
                _set_cell(session, name, r, c, shown, formula=e.formula)
            else:
                is_date = isinstance(value, (dt.date, dt.datetime))
                _set_cell(session, name, r, c, value, fmt="yyyy-mm-dd" if is_date else None)
            edited.add(e.row)
            log.append((e.row, c))
        _touch(session, name)
        _recompute(session, {name: edited}, log, name)
        df, cols = _df(session, name), _cols(session, name)
        if len(df) != n_before or list(df.columns) != names_before:
            return {"reload": True, "total_rows": len(df), "columns": list(df.columns)}
        unique = list(dict.fromkeys(log))
        if len(unique) > 20000:
            return {"reload": True, "total_rows": len(df), "columns": list(df.columns)}
        grid = session["grid"][name]
        pos = {c: p for p, c in enumerate(cols)}
        changes = []
        for i, c in unique:
            if c in pos:
                r = header_row + 1 + i
                changes.append({"row": i, "c": pos[c], "value": _clean(grid[r - 1][c - 1]),
                                "formula": _cell_formula(session, name, r, c)})
        return {"reload": False, "changes": changes, "total_rows": len(df)}


# ---------------------------------------------------------------------------
# download (built in the background, with progress)
# ---------------------------------------------------------------------------
def _sheet_has_formulas(session, name):
    try:
        with zipfile.ZipFile(session["path"]) as z, z.open(session["parts"][name]) as f:
            tail = b""
            while True:
                chunk = f.read(1 << 20)
                if not chunk:
                    return False
                blob = tail + chunk
                if b"<f>" in blob or b"<f " in blob:
                    return True
                tail = blob[-4:]
    except KeyError:
        return False


def _export_full(session, st, out):
    """Small workbook: reopen the original with openpyxl (keeps every style) and apply our changes."""
    st["phase"] = "Opening the original workbook"
    wb = openpyxl.load_workbook(session["path"])
    names = session["names"]
    for k, name in enumerate(names):
        st["done"], st["phase"] = k, f"Applying changes to {name}"
        if name in session["added"]:                                   # a sheet made from query results
            ws = wb.create_sheet(name)
            for r, row in enumerate(session["grid"][name], start=1):
                for c, v in enumerate(row, start=1):
                    if v is not None:
                        ws.cell(row=r, column=c, value=v)
            continue
        ws = wb[name]
        for r in sorted(session["deleted_rows"].get(name, []), reverse=True):
            ws.delete_rows(r, 1)
        fmts = session["wfmt"][name]
        for (r, c), v in session["writes"][name].items():
            cell = ws.cell(row=r, column=c)
            cell.value = v
            if (r, c) in fmts:
                cell.number_format = fmts[(r, c)]
    st["phase"] = "Saving"
    wb.save(out)


def _export_stream(session, st, out):
    """Big workbook: stream every sheet straight into a new file (all data + formulas, no cell styling)."""
    wbw = xlsxwriter.Workbook(out, {"constant_memory": True, "nan_inf_to_errors": True,
                                    "default_date_format": "yyyy-mm-dd"})
    f_date = wbw.add_format({"num_format": "yyyy-mm-dd"})
    f_dt = wbw.add_format({"num_format": "yyyy-mm-dd hh:mm:ss"})
    f_time = wbw.add_format({"num_format": "hh:mm:ss"})
    names = session["names"]
    for k, name in enumerate(names):
        st["done"], st["phase"] = k, f"Writing {name}"
        with session["lock"]:
            _ensure_grid(session, name)
            grid = session["grid"][name]
            formulas = {}
            if _sheet_has_formulas(session, name):                   # formulas that were already in the file
                st["phase"] = f"Reading the formulas of {name}"
                _load_orig_formulas(session, name)
                formulas.update(session["orig_formulas"].get(name, {}))
            for (r, c), v in session["writes"][name].items():
                if isinstance(v, str) and v.startswith("="):
                    formulas[(r, c)] = v
                else:
                    formulas.pop((r, c), None)                       # the user typed over a formula
            by_row = defaultdict(dict)
            for (r, c), f in formulas.items():
                by_row[r][c] = f
            ws = wbw.add_worksheet(name)
            width = max((len(r) for r in grid), default=1)
            ws.set_column(0, max(0, width - 1), 18)
            ws.freeze_panes(session["meta"][name]["header_row"], 0)
            write_str, write_num, write_bool, write_dt, write_formula = (
                ws.write_string, ws.write_number, ws.write_boolean, ws.write_datetime, ws.write_formula)
            for r0, row in enumerate(grid):
                fr = by_row.get(r0 + 1)
                for c0, v in enumerate(row):
                    if fr and (c0 + 1) in fr:
                        cached, fmt = v, None
                        if isinstance(v, dt.datetime):
                            cached, fmt = serial(v), f_dt
                        elif isinstance(v, dt.date):
                            cached, fmt = serial(v), f_date
                        if cached is None:
                            cached = 0
                        write_formula(r0, c0, fr[c0 + 1], fmt, cached)
                    elif v is None:
                        continue
                    elif isinstance(v, bool):
                        write_bool(r0, c0, v)
                    elif isinstance(v, (int, float)):
                        write_num(r0, c0, v)
                    elif isinstance(v, dt.datetime):
                        write_dt(r0, c0, v, f_dt)
                    elif isinstance(v, dt.date):
                        write_dt(r0, c0, v, f_date)
                    elif isinstance(v, dt.time):
                        write_dt(r0, c0, v, f_time)
                    else:
                        write_str(r0, c0, str(v)[:32767])
    st["phase"] = "Finishing the file"
    wbw.close()


def _export_worker(session, st):
    out = os.path.join(WORKDIR, f"export_{uuid.uuid4()}.xlsx")
    try:
        if session["mode"] == "full":
            _export_full(session, st, out)
        else:
            _export_stream(session, st, out)
        st.update(state="done", path=out, done=len(session["names"]), phase="Ready")
    except Exception as e:
        st.update(state="error", error=f"{e.__class__.__name__}: {e}")


def _start_export(session):
    with session["lock"]:
        st = session["export"]
        if st and st["state"] == "running":
            return st
        if st and st.get("path") and os.path.exists(st["path"]):
            os.remove(st["path"])
        st = {"state": "running", "done": 0, "total": len(session["names"]), "phase": "Starting", "path": None,
              "error": None, "mode": session["mode"], "started": time.time()}
        session["export"] = st
    threading.Thread(target=_export_worker, args=(session, st), daemon=True).start()
    return st


def _download_name(session):
    return (session.get("filename") or "workbook.xlsx").rsplit(".", 1)[0] + "_edited.xlsx"


XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@app.post("/api/download/start/{session_id}")
def download_start(session_id: str):
    session = _get_session(session_id)
    st = _start_export(session)
    return {k: st[k] for k in ("state", "done", "total", "phase", "mode")}


@app.get("/api/download/status/{session_id}")
def download_status(session_id: str):
    st = _get_session(session_id)["export"]
    if not st:
        return {"state": "idle"}
    return {k: st.get(k) for k in ("state", "done", "total", "phase", "error", "mode")}


@app.get("/api/download/file/{session_id}")
def download_file(session_id: str):
    session = _get_session(session_id)
    st = session["export"]
    if not st or st["state"] != "done" or not st.get("path") or not os.path.exists(st["path"]):
        raise HTTPException(409, "The file isn't ready yet.")
    return FileResponse(st["path"], media_type=XLSX_TYPE, filename=_download_name(session))


@app.get("/api/download/{session_id}")
def download_direct(session_id: str):
    """Old one-step endpoint: build (waiting if needed) and return the file."""
    session = _get_session(session_id)
    st = _start_export(session)
    while st["state"] == "running":
        time.sleep(0.3)
    if st["state"] != "done":
        raise HTTPException(500, st.get("error") or "Export failed")
    return FileResponse(st["path"], media_type=XLSX_TYPE, filename=_download_name(session))


# ---------------------------------------------------------------------------
# SQL
# ---------------------------------------------------------------------------
def _wb(session):
    return Workbench(session["names"], lambda n: _df(session, n))


def _json_cell(v):
    if isinstance(v, (bytes, bytearray)):
        return v.decode("utf8", "replace")
    return v


class SqlBody(BaseModel):
    session_id: str
    sheet_name: Optional[str] = None      # the sheet you are working on (what "this" refers to)
    sql: str
    limit: int = 1000


@app.get("/api/sql/schema/{session_id}")
def sql_schema(session_id: str):
    session = _get_session(session_id)
    with session["lock"]:
        return {"tables": _wb(session).schema()}


@app.post("/api/sql/run")
def sql_run(body: SqlBody):
    session = _get_session(body.session_id)
    try:
        res = _wb(session).run(body.sql, body.sheet_name, max(0, min(body.limit, 5000)))
    except SqlError as e:
        raise HTTPException(400, str(e))
    if res["kind"] == "select":
        res["rows"] = [[_json_cell(v) for v in r] for r in res["rows"]]
    else:
        res["shown"] = min(len(res["changes"]), 200)
        res["changes"] = res["changes"][:200]
    return res


@app.post("/api/sql/translate")
def sql_translate(body: SqlBody):
    """Can this SQL be an Excel formula? If so, hand back the builder steps plus the formula and a few values."""
    session = _get_session(body.session_id)
    if not body.sheet_name:
        return {"ok": False, "error": "Open the sheet you want to fill in first."}
    try:
        tr = _wb(session).translate(body.sheet_name, body.sql)
        outs = []
        for o in tr["outputs"]:
            probe = ApplyOperation(session_id=body.session_id, sheet_name=body.sheet_name, operation="EXPR",
                                   params={"expr": o["node"]}, output_column=o["name"])
            _, ctx, formula_fn, values = _build(probe, preview=True)
            outs.append({"name": o["name"], "node": o["node"], "formula": formula_fn(ctx["header_row"] + 1),
                         "values": [_clean(v) for v in list(values)[:6]]})
        return {"ok": True, "outputs": outs, "notes": tr["notes"], "target": tr["target"]}
    except SqlError as e:
        return {"ok": False, "error": str(e)}
    except HTTPException as e:
        return {"ok": False, "error": str(e.detail)}


def _key(v):
    """Normalise a value for matching query results to sheet rows."""
    if v is None:
        return None
    if isinstance(v, (dt.datetime, dt.date, pd.Timestamp)):
        return str(sql_value(v)).lower()
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float)) or hasattr(v, "item"):
        return float(v)
    return str(v).strip().lower()


class FillBody(BaseModel):
    session_id: str
    sheet_name: str
    sql: str
    mode: str = "by_key"                  # by_key: match a result column to a sheet column; by_row: one value per row (this.x queries)
    key_result: Optional[str] = None      # result column that identifies the row ...
    key_base: Optional[str] = None        # ... and the sheet column it must equal
    columns: List[dict] = []              # [{"result": "vehicles", "output": "vehicle_count"}]


@app.post("/api/sql/apply-values")
def sql_apply_values(body: FillBody):
    """Write query results into the current sheet as plain values (works for any SELECT, not just ones that map to formulas)."""
    session = _get_session(body.session_id)
    _check_sheet(session, body.sheet_name)
    name = body.sheet_name
    try:
        res = _wb(session).run(body.sql, name, 0)
    except SqlError as e:
        raise HTTPException(400, str(e))
    if res["kind"] != "select":
        raise HTTPException(400, "Only a SELECT can fill a column.")
    rcols = res["columns"]
    with session["lock"]:
        df, cols = _df(session, name), _cols(session, name)
        n = len(df)
        if body.mode == "by_row":
            if not res["per_row"] or "row_no" not in rcols:
                raise HTTPException(400, "Use this.<column> in the query to get one value per row.")
            ri = rcols.index("row_no")
            lookup = {int(r[ri]) - 1: r for r in res["rows"]}
            picked = [lookup.get(i) for i in range(n)]
        else:
            if body.key_result not in rcols or body.key_base not in df.columns:
                raise HTTPException(400, "Choose the result column and the sheet column that identify the same record.")
            ki = rcols.index(body.key_result)
            first = {}
            for r in res["rows"]:
                first.setdefault(_key(r[ki]), r)
            picked = [first.get(_key(v)) for v in df[body.key_base].tolist()]
        header_row, grid, written = session["meta"][name]["header_row"], session["grid"][name], []
        for spec in body.columns:
            src, out = spec.get("result"), (spec.get("output") or spec.get("result") or "").strip()
            if src not in rcols or not out:
                continue
            si = rcols.index(src)
            names = list(df.columns)
            if out in names:
                out_col = cols[names.index(out)]
            else:
                out_col = max([len(grid[0]) if grid else 0] + cols) + 1
            _set_cell(session, name, header_row, out_col, out)
            for i in range(n):
                v = from_sql_value(picked[i][si]) if picked[i] is not None else None
                _set_cell(session, name, header_row + 1 + i, out_col, v, fmt="yyyy-mm-dd" if isinstance(v, dt.date) else None)
            session["recipes"] = [r for r in session["recipes"] if (r["sheet"], r["out_col"]) != (name, out_col)]
            written.append(out)
        if not written:
            raise HTTPException(400, "Pick at least one result column to write.")
        _touch(session, name)
        df = _df(session, name)
        return {"columns": list(df.columns), "total_rows": len(df), "written": written}


class SaveSheetBody(BaseModel):
    session_id: str
    sheet_name: Optional[str] = None
    sql: str
    name: str = "Query result"


@app.post("/api/sql/save-sheet")
def sql_save_sheet(body: SaveSheetBody):
    session = _get_session(body.session_id)
    try:
        res = _wb(session).run(body.sql, body.sheet_name, 0)
    except SqlError as e:
        raise HTTPException(400, str(e))
    if res["kind"] != "select":
        raise HTTPException(400, "Only a SELECT can be saved as a sheet.")
    if len(res["rows"]) * max(1, len(res["columns"])) > 3_000_000:
        raise HTTPException(400, "That result is too big to save as a sheet (limit 3 million cells).")
    name = _add_sheet(session, body.name, res["columns"], res["rows"])
    _df(session, name)
    return {"name": name, "rows": len(res["rows"]), "sheets": _sheet_list(session)}


class CommitBody(BaseModel):
    session_id: str
    sql: str


@app.post("/api/sql/commit-update")
def sql_commit_update(body: CommitBody):
    """Apply an UPDATE for real (the same statement that was previewed as a diff)."""
    session = _get_session(body.session_id)
    try:
        res = _wb(session).run_update(body.sql)
    except SqlError as e:
        raise HTTPException(400, str(e))
    sheet = res["sheet"]
    with session["lock"]:
        header_row, cols = session["meta"][sheet]["header_row"], _cols(session, sheet)
        edited = set()
        for ch in res["changes"]:
            v = from_sql_value(ch["after"])
            _set_cell(session, sheet, header_row + 1 + ch["row"], cols[ch["col"]], v,
                      fmt="yyyy-mm-dd" if isinstance(v, dt.date) else None)
            edited.add(ch["row"])
        _touch(session, sheet)
        _recompute(session, {sheet: edited}, [], sheet)
    return {"cells": res["cells_changed"], "rows": res["rows_changed"], "sheet": sheet}


@app.get("/api/health")
def health():
    return {"status": "ok", "version": API_VERSION, "frontend": FRONTEND_VERSION, "active_sessions": len(SESSIONS)}


# ---------------------------------------------------------------------------
# The page itself. Open http://localhost:8000 - the files are sent with "no-store" so a normal refresh
# always shows the newest version (a plain static server lets the browser keep old copies).
# ---------------------------------------------------------------------------
class NoCacheStatic(StaticFiles):
    async def get_response(self, path, scope):
        response = await super().get_response(path, scope)
        response.headers["Cache-Control"] = "no-store"
        return response


FRONTEND_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "frontend")
if os.path.isdir(FRONTEND_DIR):
    app.mount("/", NoCacheStatic(directory=FRONTEND_DIR, html=True), name="ui")   # keep last: API routes win
